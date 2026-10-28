#!/usr/bin/env python3
"""wtdb.py — helper for the Wealth Tax Supporters skill.

Standard library only, except `export`, which needs openpyxl (pip install openpyxl).

Commands
  init     DIR --area NAME          create the data folder and empty CSVs
  add      DIR TABLE JSON|@FILE     add rows (JSON object, JSON list, or @file with JSON lines); IDs assigned if blank
  set      DIR TABLE ID k=v ...     update fields on one row
  validate DIR                      check rules; exit code 1 if there are errors
  schema   [--write]                check (or rewrite) the generated column blocks in skill/references/data-format.md
  dedupe   DIR                      list likely duplicate actors
  merge    DIR KEEP_ID DROP_ID      merge a duplicate actor into another (moves evidence, links, profiles)
  checkquotes DIR [--run R] [--ids E1,E2] [--sample N] [--strict]
                                    fetch source pages and check each quote appears on the page
  locate   DIR                      fill geo_unit_id from location, and rewrite location after a unit is renamed
  status   DIR                      counts, coverage grid, sweep progress, next targets
  export   DIR [--out FILE]         write the xlsx report
"""
import argparse, csv, datetime as dt, difflib, json, os, re, sys, unicodedata
from collections import Counter, defaultdict

SCHEMA = {
    "actors": ["actor_id", "name", "kind", "actor_type", "other_types", "role", "affiliation", "party", "affiliation_actor_id", "country",
               "level", "location", "geo_unit_id", "location_source", "wikidata_qid", "website", "status", "alive_checked", "prominence", "notes",
               "added_run", "updated", "refined"],
    "profiles": ["profile_id", "actor_id", "platform", "url", "source", "checked"],
    "evidence": ["evidence_id", "actor_id", "date", "stance", "grade", "evidence_type", "wt_design", "design_detail",
                 "quote_original", "quote_lang", "quote_en", "source_url", "source_title", "publisher", "archive_url",
                 "accessed", "harvest_id", "grading_note", "run_id"],
    "geo_units": ["geo_unit_id", "name", "level", "unit_type", "parent_id", "website", "sweep_status", "swept_run",
                  "swept_date", "queries_run", "hits", "notes", "code"],
    "links": ["link_id", "from_actor_id", "to_actor_id", "link_type", "evidence_id", "source_url", "note"],
    "harvest": ["harvest_id", "title", "list_type", "url", "date", "est_size", "status", "actors_added", "notes"],
    "exclusions": ["exclusion_id", "name", "reason", "detail", "source_url", "run_id"],
    "runs": ["run_id", "date", "assistant", "focus", "targets", "queries_run", "sources_opened", "candidates_seen",
             "new_actors", "new_evidence", "grade_upgrades", "saturation",
             "notes_next"],
    # Offices people hold or held. A body is a legislature chamber, council or similar; seats say how its seats are split
    # between parties (only for bodies with seat_view=y, which get a Seats tab on the site); a position ties a person to a body.
    "bodies": ["body_id", "name", "level", "kind", "seats", "as_of", "source_url", "seat_view", "notes"],
    "seats": ["seat_id", "body_id", "party", "seats", "colour", "order"],
    "positions": ["position_id", "actor_id", "body_id", "role", "party", "geo_unit_id", "start", "end", "current",
                  "member_ref", "checked", "source_url"],
    # Things to chase in this area: a name, an organisation, a document, an event or a list seen while researching, not yet in the
    # register. Leads about other areas go in data/leads.csv (leads.py). The query log shows how broadly an area was searched.
    "leads": ["lead_id", "kind", "name", "source_url", "why", "found_run", "status", "closed_run", "note"],
    "queries": ["query_id", "run_id", "date", "language", "query", "channel", "results_opened", "leads_added", "note"],
}
PREFIX = {"actors": ("A", 4), "evidence": ("E", 5), "geo_units": ("G", 4), "links": ("L", 4),
          "harvest": ("H", 4), "exclusions": ("X", 4), "runs": ("R", 3), "profiles": ("P", 5),
          "bodies": ("B", 4), "seats": ("S", 4), "positions": ("O", 5), "leads": ("D", 4), "queries": ("Q", 4)}
REQUIRED = {
    "actors": ["name", "kind", "actor_type", "country", "level", "status", "alive_checked"],
    "evidence": ["actor_id", "date", "stance", "grade", "evidence_type", "wt_design", "quote_original",
                 "quote_lang", "source_url", "grading_note"],
    "geo_units": ["name", "level", "sweep_status"],
    "links": ["from_actor_id", "to_actor_id", "link_type"],
    "profiles": ["actor_id", "platform", "url"],
    "harvest": ["title", "list_type", "status"],
    "exclusions": ["name", "reason"],
    "runs": ["date", "focus"],
    "bodies": ["name", "level", "kind"],
    "seats": ["body_id", "party", "seats"],
    "positions": ["actor_id", "body_id"],
    "leads": ["kind", "name", "source_url", "status"],
    "queries": ["query"],
}
ENUMS = {
    ("leads", "kind"): {"person", "organisation", "document", "event", "list"},
    ("leads", "status"): {"open", "used", "rejected", "duplicate"},
    ("actors", "kind"): {"person", "organisation", "institution"},
    ("actors", "actor_type"): {"politician", "party", "government_body", "academic", "economist", "think_tank",
                               "ngo_campaign", "union", "union_official", "journalist", "media_outlet",
                               "legal_tax_professional", "wealthy_individual", "business_leader", "faith",
                               "cultural_figure", "campaigner", "international_body", "funder"},
    ("actors", "level"): {"supranational", "national", "regional", "local"},
    ("actors", "status"): {"active", "former_role", "dissolved"},
    ("actors", "alive_checked"): {"y", "na"},
    ("actors", "prominence"): {"", "international", "national", "regional", "local"},
    ("bodies", "level"): {"supranational", "national", "regional", "local"},
    ("bodies", "kind"): {"chamber", "council", "other"},
    ("bodies", "seat_view"): {"", "y"},
    ("positions", "current"): {"", "y", "n"},
    ("evidence", "stance"): {"support", "retracted"},
    ("evidence", "grade"): {"A", "B", "C"},
    ("evidence", "evidence_type"): {"vote", "bill_sponsorship", "motion", "motion_signature", "open_letter",
                                    "publication", "report", "op_ed", "editorial", "speech", "statement",
                                    "interview", "social_post", "manifesto", "policy_adoption", "campaign",
                                    "coalition_membership", "testimony", "funding", "other"},
    ("evidence", "wt_design"): {"annual_net_wealth", "one_off_levy", "billionaire_minimum", "unrealised_gains",
                                "exit_tax", "partial_wealth", "defend_existing", "wealth_tax_unspecified"},
    ("geo_units", "level"): {"supranational", "national", "devolved", "region", "county", "combined_authority",
                             "council", "sub_council"},
    ("geo_units", "sweep_status"): {"pending", "in_progress", "done", "skipped"},
    ("links", "link_type"): {"member_of", "works_for", "leads", "co_author", "co_signatory", "funds", "founded",
                             "coalition_partner", "party_of"},
    ("harvest", "list_type"): {"roll_call", "cosponsors", "motion_signatories", "open_letter", "coalition_members",
                               "report_authors", "commission_evidence", "conference_motion", "other"},
    ("profiles", "platform"): {"website", "official_page", "x", "bluesky", "mastodon", "instagram", "facebook",
                               "youtube", "tiktok", "linkedin", "substack", "wikipedia"},
    ("harvest", "status"): {"queued", "in_progress", "done", "failed"},
    ("exclusions", "reason"): {"out_of_definition", "pre_2010_only", "deceased", "private_individual", "opposes",
                               "insufficient_evidence", "duplicate"},
}
PLATFORM_DOMAINS = {"x": ("x.com", "twitter.com"), "bluesky": ("bsky.app",), "instagram": ("instagram.com",),
                    "facebook": ("facebook.com", "fb.com"), "youtube": ("youtube.com", "youtu.be"),
                    "tiktok": ("tiktok.com",), "linkedin": ("linkedin.com",), "substack": ("substack.com",),
                    "wikipedia": ("wikipedia.org",)}
FORMAL_TYPES = {"vote", "bill_sponsorship", "motion", "motion_signature", "open_letter", "publication", "report",
                "manifesto", "policy_adoption", "testimony", "funding"}
GRADE_RANK = {"A": 3, "B": 2, "C": 1}
GRADE_LABEL = {"A": "explicit", "B": "named", "C": "appears to support"}
DATE_RE = re.compile(r"^\d{4}(-\d{2}(-\d{2})?)?$")
MIN_YEAR = 2010


# ---------- io ----------
def path(d, t):
    return os.path.join(d, t + ".csv")


def load(d, t):
    p = path(d, t)
    if not os.path.exists(p):
        return []
    with open(p, newline="", encoding="utf-8") as f:
        return [{k: (v or "").strip() for k, v in r.items() if k is not None} for r in csv.DictReader(f)]


def save(d, t, rows):
    p = path(d, t)
    tmp = p + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=SCHEMA[t], extrasaction="ignore", lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in SCHEMA[t]})
    os.replace(tmp, p)


def load_all(d):
    return {t: load(d, t) for t in SCHEMA}


def idcol(t):
    return SCHEMA[t][0]


def next_id(rows, t):
    pre, width = PREFIX[t]
    n = 0
    for r in rows:
        m = re.match(re.escape(pre) + r"(\d+)$", r.get(idcol(t), ""))
        if m:
            n = max(n, int(m.group(1)))
    return f"{pre}{n + 1:0{width}d}"


def norm_name(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\b(the|dr|prof|professor|sir|dame|lord|lady|baroness|cllr|councillor|mp|msp|ms|mr|mrs|rt hon)\b",
               " ", s)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def year(d):
    return int(d[:4]) if DATE_RE.match(d or "") else None


def need_dir(d):
    if not os.path.isdir(d):
        sys.exit(f"No data folder at {d}. Run: wtdb.py init {d} --area NAME")


# ---------- commands ----------
def cmd_init(a):
    os.makedirs(a.dir, exist_ok=True)
    for t in SCHEMA:
        if not os.path.exists(path(a.dir, t)):
            save(a.dir, t, [])
    prof = os.path.join(a.dir, "area_profile.md")
    if not os.path.exists(prof):
        tpl = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "references", "area-profile-template.md")
        text = open(tpl, encoding="utf-8").read() if os.path.exists(tpl) else "# Area profile — {AREA}\n"
        with open(prof, "w", encoding="utf-8") as f:
            f.write(text.replace("{AREA}", a.area))
    meta = os.path.join(a.dir, "meta.json")
    if not os.path.exists(meta):
        slug = os.path.basename(os.path.normpath(a.dir))
        with open(meta, "w", encoding="utf-8") as f:
            json.dump({"name": a.area, "iso2": slug.split("-")[0].upper(), "lead": "", "status": "in progress",
                       "boundaries": "", "notes": ""}, f, indent=2, ensure_ascii=False)
            f.write("\n")
    print(f"Initialised {a.dir} for {a.area}")


def parse_rows(arg):
    if arg.startswith("@"):
        text = open(arg[1:], encoding="utf-8").read().strip()
        if text.startswith("["):
            return json.loads(text)
        return [json.loads(l) for l in text.splitlines() if l.strip()]
    v = json.loads(arg)
    return v if isinstance(v, list) else [v]


def cmd_add(a):
    need_dir(a.dir)
    t = a.table
    if t not in SCHEMA:
        sys.exit(f"Unknown table {t}. Tables: {', '.join(SCHEMA)}")
    rows = load(a.dir, t)
    existing = {r[idcol(t)] for r in rows}
    today = dt.date.today().isoformat()
    new_ids = []
    for r in parse_rows(a.rows):
        unknown = set(r) - set(SCHEMA[t])
        if unknown:
            sys.exit(f"Unknown column(s) for {t}: {', '.join(sorted(unknown))}")
        r = {k: ("" if v is None else str(v)) for k, v in r.items()}
        if not r.get(idcol(t)):
            r[idcol(t)] = next_id(rows, t)
        elif r[idcol(t)] in existing:
            sys.exit(f"ID {r[idcol(t)]} already exists in {t}")
        if t == "actors":
            if not r.get("geo_unit_id") and r.get("location"):
                r["geo_unit_id"] = resolve_location(r["location"], geo_index(load(a.dir, "geo_units")))[0]
            r.setdefault("updated", today)
            r.setdefault("status", "active")
            if r.get("kind") in ("organisation", "institution"):
                r.setdefault("alive_checked", "na")
        if t == "evidence":
            r.setdefault("accessed", today)
            r.setdefault("stance", "support")
        if t == "geo_units":
            r.setdefault("sweep_status", "pending")
        if t == "harvest":
            r.setdefault("status", "queued")
        if t == "runs":
            r.setdefault("date", today)
        rows.append(r)
        existing.add(r[idcol(t)])
        new_ids.append(r[idcol(t)])
    save(a.dir, t, rows)
    print("Added " + ", ".join(new_ids))


def cmd_set(a):
    need_dir(a.dir)
    rows = load(a.dir, a.table)
    hit = [r for r in rows if r.get(idcol(a.table)) == a.id]
    if not hit:
        sys.exit(f"{a.id} not found in {a.table}")
    for kv in a.pairs:
        if "=" not in kv:
            sys.exit(f"Expected key=value, got {kv}")
        k, v = kv.split("=", 1)
        if k not in SCHEMA[a.table]:
            sys.exit(f"Unknown column {k} for {a.table}")
        hit[0][k] = v
    if a.table == "actors":
        hit[0]["updated"] = dt.date.today().isoformat()
    save(a.dir, a.table, rows)
    print(f"Updated {a.id}")


def validate(db):
    errs, warns = [], []
    ids = {}
    for t, rows in db.items():
        seen = set()
        for i, r in enumerate(rows, start=2):
            rid = r.get(idcol(t), "")
            where = f"{t}.csv line {i} ({rid or 'no id'})"
            if not rid:
                errs.append(f"{where}: missing {idcol(t)}")
            elif rid in seen:
                errs.append(f"{where}: duplicate ID")
            seen.add(rid)
            for c in REQUIRED[t]:
                if not r.get(c):
                    errs.append(f"{where}: missing {c}")
            for (tt, c), allowed in ENUMS.items():
                if tt == t and r.get(c) and r[c] not in allowed:
                        errs.append(f"{where}: {c}='{r.get(c)}' not allowed ({', '.join(sorted(x for x in allowed if x))})")
        ids[t] = seen
    for i, l in enumerate(db["leads"], start=2):
        where = f"leads.csv line {i} ({l.get('lead_id')})"
        if l.get("source_url") and not l["source_url"].startswith(("http://", "https://")):
            errs.append(f"{where}: source_url must start with http(s)://")
        if l.get("status") in ("used", "rejected", "duplicate") and not l.get("closed_run"):
            warns.append(f"{where}: closed without closed_run")
        for c in ("found_run", "closed_run"):
            if l.get(c) and l[c] not in ids["runs"]:
                warns.append(f"{where}: {c} {l[c]} is not in runs.csv (add the run row before the run ends)")
    for i, q in enumerate(db["queries"], start=2):
        if q.get("run_id") and q["run_id"] not in ids["runs"]:
            warns.append(f"queries.csv line {i} ({q.get('query_id')}): run_id {q['run_id']} is not in runs.csv")
    actors = {r["actor_id"]: r for r in db["actors"]}
    ev_by_actor = defaultdict(list)
    for i, e in enumerate(db["evidence"], start=2):
        where = f"evidence.csv line {i} ({e.get('evidence_id')})"
        ev_by_actor[e.get("actor_id")].append(e)
        if e.get("actor_id") and e["actor_id"] not in actors:
            errs.append(f"{where}: actor_id {e['actor_id']} does not exist")
        d = e.get("date", "")
        if d and not DATE_RE.match(d):
            errs.append(f"{where}: date '{d}' must be YYYY-MM-DD, YYYY-MM or YYYY")
        elif d and year(d) < MIN_YEAR:
            errs.append(f"{where}: date {d} is before {MIN_YEAR}")
        elif d and d > dt.date.today().isoformat():
            errs.append(f"{where}: date {d} is in the future")
        u = e.get("source_url", "")
        if u and not u.startswith(("http://", "https://")):
            errs.append(f"{where}: source_url must start with http(s)://")
        if e.get("quote_lang") and e["quote_lang"].lower() not in ("en", "eng") and not e.get("quote_en"):
            errs.append(f"{where}: quote_en (translation) needed for quote_lang={e['quote_lang']}")
        if e.get("harvest_id") and e["harvest_id"] not in ids["harvest"]:
            errs.append(f"{where}: harvest_id {e['harvest_id']} does not exist")
        if e.get("grade") == "A" and e.get("evidence_type") not in FORMAL_TYPES and \
                e.get("wt_design") == "wealth_tax_unspecified" and not e.get("design_detail"):
            warns.append(f"{where}: grade A on a {e.get('evidence_type')} with no design or detail — should this be B?")
        if e.get("grade") == "A" and e.get("evidence_type") not in FORMAL_TYPES and not e.get("design_detail") \
                and e.get("wt_design") != "wealth_tax_unspecified":
            warns.append(f"{where}: grade A {e.get('evidence_type')} — add design_detail (rate, threshold or bill)")
        if len(e.get("quote_original", "")) < 15 and e.get("evidence_type") != "vote":  # a vote quotes the member's name on the list
            warns.append(f"{where}: very short quote")
    geo = {g["geo_unit_id"] for g in db["geo_units"]}
    gidx = geo_index(db["geo_units"])
    names = defaultdict(list)
    qids = defaultdict(list)
    for i, r in enumerate(db["actors"], start=2):
        where = f"actors.csv line {i} ({r.get('actor_id')} {r.get('name')})"
        if r.get("actor_id") not in ev_by_actor:
            errs.append(f"{where}: no evidence — every actor needs at least one evidence row")
        if r.get("kind") == "person" and r.get("alive_checked") != "y":
            errs.append(f"{where}: person must have alive_checked=y")
        if r.get("kind") != "person" and r.get("alive_checked") == "y" and r.get("status") == "dissolved":
            warns.append(f"{where}: dissolved organisation marked alive_checked=y")
        uid, depth, rest = resolve_location(r.get("location", ""), gidx)
        gu = r.get("geo_unit_id", "")
        if gu and gu not in geo:
            errs.append(f"{where}: geo_unit_id {gu} is not in geo_units.csv")
        elif uid and gu != uid:
            errs.append(f"{where}: location '{r.get('location')}' resolves to {uid or 'nothing'} but geo_unit_id is "
                        f"{gu or 'empty'}. Run `wtdb.py locate` (it fills empty IDs and rewrites locations after a "
                        f"unit is renamed), or fix whichever one is wrong")
        if not r.get("location"):
            warns.append(f"{where}: no location (use at least the country name)")
        elif gidx[2] and not uid:
            errs.append(f"{where}: location must start with the country root unit "
                        f"({', '.join(gidx[0][u]['name'] for u in gidx[2].values())})")
        elif r.get("level") == "local" and depth <= 1 and not rest:
            warns.append(f"{where}: local actor located only at country level")
        if re.search(r"\balive\b|death record|no death|wikidata (id|match)|no wikidata", r.get("notes", ""), re.I):
            warns.append(f"{where}: notes narrate the alive or Wikidata check; alive_checked records it, so remove that from notes")
        for ot in type_list(r)[1:]:
            if ot not in ENUMS[("actors", "actor_type")]:
                errs.append(f"{where}: other_types has unknown type '{ot}'")
        if "politician" in type_list(r) and not r.get("party") and r.get("level") != "local":
            warns.append(f"{where}: politician without party (use the party's common name, or Independent; "
                         f"local officials may leave it blank when it is not known)")
        if r.get("affiliation_actor_id") and r["affiliation_actor_id"] not in actors:
            errs.append(f"{where}: affiliation_actor_id {r['affiliation_actor_id']} does not exist")
        if r.get("wikidata_qid") and not re.match(r"^Q\d+$", r["wikidata_qid"]):
            errs.append(f"{where}: wikidata_qid must look like Q12345")
        if r.get("location_source") and not r["location_source"].startswith(("https://", "http://")):
            errs.append(f"{where}: location_source must be a link to the page that shows where they are")
        if r.get("refined") and not re.match(r"^\d{4}-\d{2}-\d{2}$", r["refined"]):
            errs.append(f"{where}: refined must be a date (YYYY-MM-DD)")
        names[(norm_name(r.get("name", "")), r.get("kind"))].append(r.get("actor_id"))
        if r.get("wikidata_qid"):
            qids[r["wikidata_qid"]].append(r.get("actor_id"))
    for (n, k), v in names.items():
        if len(v) > 1:
            errs.append(f"actors: same name '{n}' ({k}) on {', '.join(v)} — merge, or add a distinguishing note "
                        f"and different names if they really are different people")
    for q, v in qids.items():
        if len(v) > 1:
            errs.append(f"actors: wikidata {q} used by {', '.join(v)} — duplicate")
    for i, l in enumerate(db["links"], start=2):
        for c in ("from_actor_id", "to_actor_id"):
            if l.get(c) and l[c] not in actors:
                errs.append(f"links.csv line {i}: {c} {l[c]} does not exist")
        if l.get("evidence_id") and l["evidence_id"] not in ids["evidence"]:
            errs.append(f"links.csv line {i}: evidence_id {l['evidence_id']} does not exist")
    seen_profiles = set()
    for i, p in enumerate(db["profiles"], start=2):
        where = f"profiles.csv line {i}"
        a = actors.get(p.get("actor_id"))
        if not a:
            errs.append(f"{where}: actor_id {p.get('actor_id')} does not exist")
        url = p.get("url", "")
        m = re.match(r"^https://([^/\s]+)(/\S*)?$", url)
        if not m:
            errs.append(f"{where}: url must be an https link (no mailto:, tel: or bare handles)")
        else:
            host = m.group(1).lower().split(":")[0]
            doms = PLATFORM_DOMAINS.get(p.get("platform"))
            if doms and not any(host == d or host.endswith("." + d) for d in doms):
                errs.append(f"{where}: {url} does not look like a {p.get('platform')} link")
        key = (p.get("actor_id"), url.rstrip("/").lower())
        if key in seen_profiles:
            warns.append(f"{where}: duplicate link for {p.get('actor_id')}")
        seen_profiles.add(key)
        if a and set(type_list(a)) == {"wealthy_individual"}:
            warns.append(f"{where}: {a.get('name')} is only a wealthy individual; link only official or professional pages of public figures")
    bodies = {b["body_id"]: b for b in db["bodies"]}
    for i, b in enumerate(db["bodies"], start=2):
        where = f"bodies.csv line {i} ({b.get('body_id')} {b.get('name')})"
        if b.get("seat_view") == "y":
            for c in ("seats", "as_of", "source_url"):
                if not b.get(c):
                    errs.append(f"{where}: seat_view=y needs {c}")
        if b.get("seats") and not b["seats"].isdigit():
            errs.append(f"{where}: seats must be a whole number")
        if b.get("as_of") and not DATE_RE.match(b["as_of"]):
            errs.append(f"{where}: as_of must be a date")
    split = defaultdict(int)
    for i, s in enumerate(db["seats"], start=2):
        where = f"seats.csv line {i} ({s.get('seat_id')})"
        if s.get("body_id") not in bodies:
            errs.append(f"{where}: body_id {s.get('body_id')} does not exist")
        if not s.get("seats", "").isdigit():
            errs.append(f"{where}: seats must be a whole number")
        else:
            split[s.get("body_id")] += int(s["seats"])
    for bid, n in split.items():
        b = bodies.get(bid)
        if b and b.get("seats", "").isdigit() and n != int(b["seats"]):
            warns.append(f"seats: {b['name']} has {b['seats']} seats but its parties add up to {n} (add a 'Vacant' row)")
    today_s = dt.date.today().isoformat()
    cur = defaultdict(list)
    for i, o in enumerate(db["positions"], start=2):
        where = f"positions.csv line {i} ({o.get('position_id')})"
        a = actors.get(o.get("actor_id"))
        if not a:
            errs.append(f"{where}: actor_id {o.get('actor_id')} does not exist")
        if o.get("body_id") not in bodies:
            errs.append(f"{where}: body_id {o.get('body_id')} does not exist")
        for c in ("start", "end", "checked"):
            if o.get(c) and not DATE_RE.match(o[c]):
                errs.append(f"{where}: {c} must be a date (YYYY-MM-DD, YYYY-MM or YYYY)")
        if o.get("geo_unit_id") and o["geo_unit_id"] not in geo:
            errs.append(f"{where}: geo_unit_id {o['geo_unit_id']} does not exist")
        if o.get("current") == "y":
            if o.get("end"):
                errs.append(f"{where}: current=y but it has an end date")
            if not o.get("checked"):
                errs.append(f"{where}: current=y needs the date it was checked against the body's own list")
            cur[o.get("actor_id")].append(o)
        if o.get("current") == "n" and not o.get("checked"):
            warns.append(f"{where}: current=n without a checked date")
        if o.get("end") and o["end"] > today_s:
            errs.append(f"{where}: end date is in the future")
    for aid, os_ in cur.items():
        a = actors.get(aid)
        if a and a.get("status") == "former_role" and any(bodies.get(o["body_id"], {}).get("kind") == "chamber" for o in os_):
            warns.append(f"actors: {a.get('name')} is former_role but holds a current seat in a chamber")
    codes = defaultdict(list)
    for i, g in enumerate(db["geo_units"], start=2):
        if g.get("code"):
            codes[g["code"]].append(g["geo_unit_id"])
        if g.get("parent_id") and g["parent_id"] not in geo:
            errs.append(f"geo_units.csv line {i}: parent_id {g['parent_id']} does not exist")
        if g.get("sweep_status") == "done" and not g.get("swept_date"):
            warns.append(f"geo_units.csv line {i} ({g.get('name')}): done without swept_date")
    for c, v in codes.items():
        if len(v) > 1:
            errs.append(f"geo_units: code {c} used by {', '.join(v)}")
    return errs, warns


# Licences that let anyone republish map shapes. A map from a source with no stated licence is not published.
OPEN_LICENCES = ("cc0", "cc-by", "cc by", "odbl", "pddl", "public domain", "ogl", "open government licence", "dl-de",
                 "datenlizenz deutschland", "godl", "licence ouverte", "etalab", "nlod", "use without restriction", "libre uso")
LOCAL_LEVELS = {"county", "combined_authority", "council", "sub_council"}
LEGISLATOR = re.compile(r"(?i)\b(member of (the )?(congress|parliament|bundestag|lok sabha|riksdag|house of (commons|representatives|delegates)|assembl|national assembly)"
                        r"|MP\b|MdB\b|député|diputad|deputy\b(?!\s*,|\s+(chair|head|leader|director|general|secretary|mayor|prime|minister|governor))|senator|sénat|congress(wo)?man|representative\b)")


def check_boundaries(d, db):
    """Errors and warnings for boundaries.json: its licence and code scheme, and shapes that match no geo unit."""
    errs, warns = [], []
    try:
        meta = json.load(open(os.path.join(d, "meta.json"), encoding="utf-8"))
    except (OSError, ValueError):
        return errs, warns
    units = [g for g in db["geo_units"] if g.get("level") in LOCAL_LEVELS]
    if db["geo_units"] and not units:
        errs.append("geo_units: the register stops above the local tier. Load the official list of local units with "
                    "their codes (SKILL.md 6.7) before research; it is part of every seed pass")
    root = next((g["name"] for g in db["geo_units"] if g.get("unit_type") == "country"), "")
    national = re.compile(meta.get("national_seat_roles") or "(?!)")  # roles whose seat is the whole country (Colombia's Senate, indigenous and opposition seats)
    seated = [a["name"] for a in db["actors"] if root and a.get("location") == root and LEGISLATOR.search(a.get("role", "")) and not national.search(a.get("role", ""))]
    if units and seated:
        warns.append(f"actors: {len(seated)} legislators are placed at the country root, for example {', '.join(seated[:3])}. "
                     f"Seats are places: put each in their constituency (SKILL.md 6.7)")
    bfile = meta.get("boundaries")
    if not bfile or not os.path.exists(os.path.join(d, bfile)):
        warns.append("boundaries: no map. Add a local-tier map from an open source (SKILL.md 5 and 6.7); a missing map is not a blocker: "
                     "if none exists, say which sources you checked in meta.json notes and open a follow-up issue for the map")
        return errs, warns
    b = json.load(open(os.path.join(d, bfile), encoding="utf-8"))
    lic = (b.get("licence") or "").strip()
    if not lic or lic.lower() in ("unknown", "none", "not stated", "unclear"):
        errs.append("boundaries: no licence. Use shapes from a source with an open licence and record it in "
                    "\"licence\" and \"licence_url\" (SKILL.md 5, boundaries.json); a map without one is not published")
    elif not any(k in lic.lower() or k in lic.lower().replace("-", " ") for k in OPEN_LICENCES):
        warns.append(f"boundaries: licence '{lic}' is not one of the known open licences; check that it allows "
                     f"republishing before merging")
    if lic and not b.get("licence_url", "").startswith("https://"):
        errs.append("boundaries: licence_url must link to the licence text (https)")
    if not b.get("code_scheme"):
        errs.append("boundaries: no code_scheme. Name the official code system of the units (GSS, FIPS, AGS, "
                    "kommunkod, LGD ...); codes must never be invented")
    by_code = {g["code"] for g in db["geo_units"] if g.get("code")}
    by_name = {_bare(g["name"]) for g in db["geo_units"]}
    miss = [u["name"] for u in b.get("units", [])
            if not ((u.get("code") and u["code"] in by_code) or (not by_code and _bare(u["name"]) in by_name))]
    if miss:
        warns.append(f"boundaries: {len(miss)} map shapes match no geo unit (by code{'' if by_code else ' or name'}), "
                     f"for example {', '.join(miss[:5])}")
    if not by_code:
        warns.append("boundaries: geo_units has no codes; give each mapped unit its official code (column code)")
    return errs, warns


def validate_supranational(d):
    """A body above the countries (data/eu): only bodies, seats and positions. Its supporters are actors in the country registers, referred
    to as <country folder>-<actor id>; check each one exists there, and that a body's seat rows add up to its seats."""
    errs, warns = [], []
    db = {t: load(d, t) for t in ("bodies", "seats", "positions", "actors", "evidence", "exclusions", "runs")}  # actors, evidence: only a body whose members are governments (data/un) holds its own
    for t, rows in db.items():
        seen = set()
        for i, r in enumerate(rows, start=2):
            rid = r.get(idcol(t), "")
            where = f"{t}.csv line {i} ({rid or 'no id'})"
            if not rid or rid in seen:
                errs.append(f"{where}: missing or duplicate {idcol(t)}")
            seen.add(rid)
            errs += [f"{where}: missing {c}" for c in REQUIRED[t] if not r.get(c)]
            for (tt, c), allowed in ENUMS.items():
                if tt == t and r.get(c) and r[c] not in allowed:
                    errs.append(f"{where}: {c}='{r[c]}' not allowed")
    bodies = {b["body_id"]: b for b in db["bodies"]}
    for b in db["bodies"]:
        total = sum(int(x["seats"] or 0) for x in db["seats"] if x["body_id"] == b["body_id"])
        if b.get("seat_view") == "y" and total != int(b["seats"] or 0):
            errs.append(f"bodies.csv ({b['body_id']}): seats.csv rows add up to {total}, not {b['seats']}")
    own = {x["actor_id"] for x in db["actors"]}
    for e in db["evidence"]:
        if e["actor_id"] not in own:
            errs.append(f"evidence.csv ({e['evidence_id']}): actor {e['actor_id']} is not in this folder's actors.csv")
    for a in db["actors"]:
        if not any(e["actor_id"] == a["actor_id"] for e in db["evidence"]):
            errs.append(f"actors.csv ({a['actor_id']}): no evidence row")
    root = os.path.dirname(os.path.abspath(d))
    actors = {}
    for o in db["positions"]:
        folder, _, aid = o["actor_id"].rpartition("-")
        if folder not in actors:
            f = os.path.join(root, folder, "actors.csv")
            actors[folder] = {x["actor_id"] for x in load(os.path.join(root, folder), "actors")} if os.path.exists(f) else set()
        if aid not in actors[folder]:
            errs.append(f"positions.csv ({o['position_id']}): actor {o['actor_id']} is not in data/{folder or '?'}/actors.csv")
        if o["body_id"] not in bodies:
            errs.append(f"positions.csv ({o['position_id']}): body_id {o['body_id']} does not exist")
    return errs, warns


def cmd_validate(a):
    need_dir(a.dir)
    meta = os.path.join(a.dir, "meta.json")
    if os.path.exists(meta) and json.load(open(meta, encoding="utf-8")).get("supranational"):
        errs, warns = validate_supranational(a.dir)
        for w in warns:
            print("WARN  " + w)
        for e in errs:
            print("ERROR " + e)
        print(f"\n{len(errs)} error(s), {len(warns)} warning(s) (supranational body: bodies, seats and positions only)")
        sys.exit(1 if errs else 0)
    db = load_all(a.dir)
    errs, warns = validate(db)
    if os.path.exists(meta):
        warns += check_folder(a.dir, json.load(open(meta, encoding="utf-8")))
    have = {norm_name(r["name"]) for r in db["actors"]}
    for x in db["exclusions"]:
        if norm_name(x["name"]) in have:
            warns.append(f"exclusions.csv {x['exclusion_id']}: '{x['name']}' is also an actor; remove one of the rows")
    be, bw = check_boundaries(a.dir, db)
    errs += be
    warns += bw
    for w in warns:
        print("WARN  " + w)
    for e in errs:
        print("ERROR " + e)
    print(f"\n{len(errs)} error(s), {len(warns)} warning(s)")
    sys.exit(1 if errs else 0)


def check_folder(d, meta):
    """Warning when the folder name is not the lowercase ISO code (or <iso2>-<region>), unless the code is taken by another folder."""
    slug = os.path.basename(os.path.abspath(d))
    iso = (meta.get("iso2") or "").lower()
    if not iso or slug == iso or slug.startswith(iso + "-"):
        return []
    parent = os.path.dirname(os.path.abspath(d))
    for other in os.listdir(parent):
        mp = os.path.join(parent, other, "meta.json")
        if other != slug and os.path.exists(mp) and (json.load(open(mp, encoding="utf-8")).get("iso2") or "").lower() == iso:
            return []
    return [f"folder: '{slug}' is not the lowercase ISO code '{iso}' (or {iso}-<region>). Use the code, or the country's name only when another folder already holds the code (SKILL.md 5)"]


# ---------- the generated column blocks in data-format.md ----------
DOC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "references", "data-format.md")


def schema_block(t):
    pre, width = PREFIX[t]
    lines = [f"`{', '.join(SCHEMA[t])}`",
             f"- IDs look like `{pre}{1:0{width}d}`. Required: " + (", ".join(f"`{c}`" for c in REQUIRED.get(t, [])) or "none") + "."]
    for (tt, c), vals in ENUMS.items():
        if tt == t:
            lines.append(f"- `{c}` is one of: " + " · ".join(sorted(v for v in vals if v)) + (" (or blank)" if "" in vals else "") + ".")
    return "\n".join(lines)


def cmd_schema(a):
    text = open(DOC, encoding="utf-8", newline="").read()
    new, n = text, 0
    for t in SCHEMA:
        pat = re.compile(r"(<!-- schema:%s -->\n).*?(\n<!-- /schema:%s -->)" % (t, t), re.S)
        if not pat.search(new):
            sys.exit(f"data-format.md has no schema block for {t}")
        new = pat.sub(lambda m: m.group(1) + schema_block(t) + m.group(2), new)
        n += 1
    if new == text:
        print(f"data-format.md is up to date ({n} generated blocks)")
        return
    if a.write:
        open(DOC, "w", encoding="utf-8", newline="").write(new)
        print("data-format.md rewritten from SCHEMA, PREFIX, REQUIRED and ENUMS")
        return
    print("data-format.md is out of date with wtdb.py: run `python skill/scripts/wtdb.py schema --write` and commit the result")
    sys.exit(1)


def cmd_dedupe(a):
    need_dir(a.dir)
    actors = load(a.dir, "actors")
    pairs = []
    for i in range(len(actors)):
        for j in range(i + 1, len(actors)):
            x, y = actors[i], actors[j]
            if x["kind"] != y["kind"]:
                continue
            nx, ny = norm_name(x["name"]), norm_name(y["name"])
            r = difflib.SequenceMatcher(None, nx, ny).ratio()
            px, py = nx.split(), ny.split()
            same_last = (x["kind"] == "person" and len(px) > 1 and len(py) > 1 and px[-1] == py[-1]
                         and px[0][0] == py[0][0])
            contains = x["kind"] != "person" and len(min(nx, ny, key=len)) > 4 and (nx in ny or ny in nx)
            if r >= 0.86 or same_last or contains:
                pairs.append((r, x, y))
    pairs.sort(key=lambda p: -p[0])
    for r, x, y in pairs:
        print(f"{r:.2f}  {x['actor_id']} {x['name']} ({x['actor_type']}, {x.get('affiliation','')})  <->  "
              f"{y['actor_id']} {y['name']} ({y['actor_type']}, {y.get('affiliation','')})")
    print(f"\n{len(pairs)} possible duplicate pair(s). Check each; merge true duplicates (keep the older ID).")


def derive(db):
    ev_by_actor = defaultdict(list)
    for e in db["evidence"]:
        ev_by_actor[e["actor_id"]].append(e)
    out = []
    for r in db["actors"]:
        evs = sorted(ev_by_actor.get(r["actor_id"], []), key=lambda e: e.get("date", ""))
        sup = [e for e in evs if e.get("stance") == "support"]
        latest = evs[-1] if evs else {}
        best = max((e.get("grade", "C") for e in sup), key=lambda g: GRADE_RANK.get(g, 0), default="")
        strongest = max(sup, key=lambda e: (GRADE_RANK.get(e.get("grade"), 0), e.get("date", "")), default={})
        out.append(dict(r,
                        best_grade=best,
                        support_label=GRADE_LABEL.get(best, ""),
                        support_status="retracted" if latest.get("stance") == "retracted" else "supporter",
                        evidence_count=str(len(sup)),
                        designs="; ".join(sorted({e.get("wt_design") for e in sup if e.get("wt_design")})),
                        evidence_types="; ".join(sorted({e.get("evidence_type") for e in sup})),
                        first_date=sup[0]["date"] if sup else "",
                        latest_date=sup[-1]["date"] if sup else "",
                        strongest_quote=strongest.get("quote_en") or strongest.get("quote_original", ""),
                        strongest_source=strongest.get("source_url", "")))
    return out


def _bare(n):
    return re.sub(r"\s*\(.*$", "", n or "").strip().lower()


def geo_index(geo):
    """Index the geography register: id lookup, children by bare name, and the country root unit(s)."""
    by_id = {g["geo_unit_id"]: g for g in geo}
    kids, roots = defaultdict(dict), {}
    for g in geo:
        if g.get("parent_id") and g["parent_id"] in by_id:
            kids[g["parent_id"]].setdefault(_bare(g["name"]), g["geo_unit_id"])
        elif g.get("unit_type") == "country":
            roots[_bare(g["name"])] = g["geo_unit_id"]
    return by_id, kids, roots


def resolve_location(location, idx):
    """Walk an actor's 'Country > Region > Place' path down the register.
    Returns (unit_id or '', number of segments matched, the segments left over as plain place text)."""
    by_id, kids, roots = idx
    parts = [p.strip() for p in (location or "").split(">") if p.strip()]
    uid = roots.get(_bare(parts[0])) if parts else None
    if not uid:
        return "", 0, parts
    depth = 1
    for seg in parts[1:]:
        nxt = kids.get(uid, {}).get(_bare(seg))
        if not nxt:
            break
        uid, depth = nxt, depth + 1
    return uid, depth, parts[depth:]


def type_list(r):
    """Primary type first, then any extra types (other_types is semicolon-separated)."""
    extra = [x.strip() for x in (r.get("other_types") or "").split(";") if x.strip()]
    return [r["actor_type"]] + [x for x in extra if x != r["actor_type"]]


def coverage(db, derived):
    grid = defaultdict(Counter)
    for r in derived:
        if r["support_status"] == "supporter":
            for ty in type_list(r):
                grid[ty][r["level"]] += 1
    return grid


def sweep_progress(db):
    """Count leaf units (areas nothing else sits under) by sweep status, grouped by their region."""
    geo = db["geo_units"]
    by_id = {g["geo_unit_id"]: g for g in geo}
    parents = {g["parent_id"] for g in geo if g.get("parent_id")}
    prog = defaultdict(Counter)
    for g in geo:
        if g["level"] in ("national", "supranational") or g["geo_unit_id"] in parents:
            continue
        p, seen, group = g, set(), "(top level)"
        while p.get("parent_id") and p["parent_id"] in by_id and p["parent_id"] not in seen:
            seen.add(p["parent_id"])
            p = by_id[p["parent_id"]]
            if p["level"] in ("region", "devolved"):
                group = p["name"]
                break
        prog[group][g["sweep_status"]] += 1
    return prog


def unit_names(uid, by_id):
    names, seen = [], set()
    while uid and uid in by_id and uid not in seen:
        seen.add(uid); names.append(by_id[uid]["name"]); uid = by_id[uid].get("parent_id")
    return names[::-1]


def cmd_locate(a):
    """geo_unit_id is the actor's place; location is its readable path plus any finer place text.
    Fill empty geo_unit_ids from location, and rewrite the unit part of location from the register when a unit was
    renamed or moved (the place text after it is kept)."""
    need_dir(a.dir)
    geo = load(a.dir, "geo_units")
    gidx = geo_index(geo)
    by_id = gidx[0]
    rows = load(a.dir, "actors")
    filled = rewritten = 0
    for r in rows:
        uid, depth, rest = resolve_location(r.get("location", ""), gidx)
        gu = r.get("geo_unit_id", "")
        if not gu:
            if uid:
                r["geo_unit_id"] = uid; filled += 1
            continue
        if gu in by_id and gu != uid:
            names = unit_names(gu, by_id)
            parts = [p.strip() for p in r.get("location", "").split(">") if p.strip()]
            r["location"] = " > ".join(names + parts[len(names):]); rewritten += 1
    save(a.dir, "actors", rows)
    print(f"geo_unit_id filled for {filled} actors; location rewritten for {rewritten}")


def cmd_status(a):
    need_dir(a.dir)
    db = load_all(a.dir)
    d = derive(db)
    sup = [r for r in d if r["support_status"] == "supporter"]
    print(f"Actors: {len(sup)} supporters, {len(d) - len(sup)} former supporters | "
          f"evidence rows: {len(db['evidence'])} | links: {len(db['links'])} | exclusions: {len(db['exclusions'])}")
    g = Counter(r["best_grade"] for r in sup)
    print(f"Grades: A explicit {g['A']}, B named {g['B']}, C appears to support {g['C']}")
    print("Kinds: " + ", ".join(f"{k} {v}" for k, v in Counter(r['kind'] for r in sup).most_common()))
    levels = ["supranational", "national", "regional", "local"]
    grid = coverage(db, d)
    types = sorted(ENUMS[("actors", "actor_type")])
    print("\nCoverage grid (supporters by type × level):")
    print(f"  {'actor_type':24}" + "".join(f"{l[:9]:>10}" for l in levels) + f"{'total':>8}")
    for t in types:
        row = grid.get(t, Counter())
        print(f"  {t:24}" + "".join(f"{row[l]:>10}" for l in levels) + f"{sum(row.values()):>8}")
    prog = sweep_progress(db)
    if prog:
        print("\nLocal sweep progress (units by status, grouped by region):")
        for reg, c in sorted(prog.items()):
            tot = sum(c.values())
            print(f"  {reg:32} done {c['done']:>4}/{tot:<4} pending {c['pending']:>4}  in_progress {c['in_progress']:>3}"
                  f"  skipped {c['skipped']:>3}")
    else:
        print("\nLocal sweep: no geography units loaded yet.")
    lq = Counter(l["status"] for l in db["leads"])
    print(f"\nLeads in this area: open {lq['open']}, used {lq['used']}, rejected {lq['rejected']}, duplicate {lq['duplicate']} | "
          f"queries logged: {len(db['queries'])} in {len({q['language'] for q in db['queries'] if q.get('language')})} language(s)")
    hq = Counter(h["status"] for h in db["harvest"])
    print(f"\nHarvest lists: queued {hq['queued']}, in_progress {hq['in_progress']}, done {hq['done']}, "
          f"failed {hq['failed']}")
    if db["runs"]:
        last = db["runs"][-1]
        print(f"\nLast run {last['run_id']} ({last['date']}): {last['focus']} → next: {last.get('notes_next','')}")
    print("\nSuggested next targets:")
    sugg = []
    ol = [l for l in db["leads"] if l["status"] == "open"]
    if ol:
        sugg.append(f"follow the {len(ol)} open lead(s) first (SKILL.md 6.3a): " + "; ".join(f"{l['lead_id']} {l['name']}" for l in ol[:5]))
    q = [h for h in db["harvest"] if h["status"] in ("queued", "in_progress")]
    if q:
        sugg.append(f"process harvest lists: " + "; ".join(f"{h['harvest_id']} {h['title']}" for h in q[:5]))
    thin = [t for t in types if sum(grid.get(t, Counter()).values()) < 3]
    if thin:
        sugg.append("thin actor types (cover these layers first): " + ", ".join(thin))
    if g["C"]:
        sugg.append(f"upgrade {g['C']} C-graded actors")
    pend = [(reg, c) for reg, c in prog.items() if c["pending"] or c["in_progress"]]
    if pend:
        pend.sort(key=lambda x: (-(x[1]["in_progress"]), -x[1]["done"]))
        sugg.append(f"regional and local units remain ({pend[0][0]}: {pend[0][1]['pending']} pending). "
                    f"Sweep them only if a sample showed they pay (SKILL.md 6.5)")
    elif not db["geo_units"]:
        sugg.append("load the geography register to the depth you will use (SKILL.md 6.7)")
    for s in sugg or ["all queues empty — run an update sweep for new evidence since the last run"]:
        print("  - " + s)


def cmd_export(a):
    need_dir(a.dir)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("export needs openpyxl: pip install openpyxl (the CSVs are the database and can be shared as they are)")
    db = load_all(a.dir)
    errs, _ = validate(db)
    if errs:
        print(f"Note: {len(errs)} validation error(s) — run validate. Exporting anyway.")
    d = derive(db)
    slug = os.path.basename(os.path.normpath(a.dir))
    out = a.out or os.path.join(a.dir, f"{slug}_wealth_tax_supporters.xlsx")
    actors = {r["actor_id"]: r for r in db["actors"]}
    gidx = geo_index(db["geo_units"])
    geo_count = Counter()  # supporters in each unit, counting everyone located in it or below it
    for r in d:
        if r["support_status"] != "supporter":
            continue
        uid = resolve_location(r.get("location", ""), gidx)[0]
        seen = set()
        while uid and uid in gidx[0] and uid not in seen:
            seen.add(uid); geo_count[uid] += 1; uid = gidx[0][uid].get("parent_id", "")
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF")

    def sheet(title, header, rows, widths=None, url_cols=()):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, h in enumerate(header, start=1):
            w = (widths or {}).get(h)
            if w is None:
                vals = [len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r)]
                w = min(max(vals) + 2, 60)
            ws.column_dimensions[get_column_letter(i)].width = w
        for h in url_cols:
            if h in header:
                ci = header.index(h) + 1
                for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                    for c in row:
                        if c.value and str(c.value).startswith("http"):
                            c.hyperlink = str(c.value)
                            c.font = Font(color="1F5FBF", underline="single")
        ws.freeze_panes = "B2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        return ws

    # Overview
    ws = wb.active
    ws.title = "Overview"
    sup = [r for r in d if r["support_status"] == "supporter"]
    g = Counter(r["best_grade"] for r in sup)
    lines = [
        ("Wealth tax supporter map", ""),
        ("Area folder", slug),
        ("Exported", dt.date.today().isoformat()),
        ("Supporters", len(sup)),
        ("Former supporters (retracted)", len(d) - len(sup)),
        ("A — explicit", g["A"]), ("B — named", g["B"]), ("C — appears to support", g["C"]),
        ("Evidence rows", len(db["evidence"])),
        ("Geography units done / total",
         f"{sum(1 for x in db['geo_units'] if x['sweep_status'] == 'done')} / {len(db['geo_units'])}"),
        ("Harvest lists done / total",
         f"{sum(1 for x in db['harvest'] if x['status'] == 'done')} / {len(db['harvest'])}"),
        ("Runs", len(db["runs"])),
        ("", ""),
        ("Definition", "Actual wealth taxation since 2010: annual net wealth tax, one-off wealth levy, billionaire "
                       "minimum tax as % of wealth, taxing unrealised gains of the very wealthy, exit taxes, broad "
                       "partial wealth taxes, defending an existing wealth tax, or an unspecified wealth tax. "
                       "Excludes mansion/property/council tax, land value tax, CGT, inheritance, income, windfall, "
                       "corporate taxes and generic 'tax the rich'."),
        ("Grade A — explicit", "Specific design (rate, threshold, bill) or a formal act: vote, sponsorship, signed "
                               "motion or letter, policy adoption, publication, testimony."),
        ("Grade B — named", "Supports 'a wealth tax' by name without specifics."),
        ("Grade C — appears to support", "Ambiguous language that probably means wealth taxation; flagged."),
        ("No ranking", "Filter the Supporters tab by grade, type, level, prominence and dates."),
    ]
    for k, v in lines:
        ws.append([k, v])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=14):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    grid = coverage(db, d)
    levels = ["supranational", "national", "regional", "local"]
    ws.append([])
    ws.append(["Coverage: supporters by type × level"] + levels + ["total"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for t in sorted(ENUMS[("actors", "actor_type")]):
        row = grid.get(t, Counter())
        ws.append([t] + [row[l] for l in levels] + [sum(row.values())])

    cols = ["actor_id", "name", "kind", "actor_type", "other_types", "role", "affiliation", "party", "country", "level", "location",
            "prominence", "best_grade", "support_label", "designs", "evidence_types", "evidence_count", "first_date",
            "latest_date", "strongest_quote", "strongest_source", "wikidata_qid", "website", "status", "notes"]
    geo_name = {x["geo_unit_id"]: x["name"] for x in db["geo_units"]}
    order = sorted(sup, key=lambda r: (-GRADE_RANK.get(r["best_grade"], 0), r["actor_type"], r["name"]))
    sheet("Supporters", cols,
          [[r.get(c, "") for c in cols] for r in order],
          widths={"strongest_quote": 70, "notes": 50, "strongest_source": 40}, url_cols=("strongest_source", "website"))
    ecols = ["evidence_id", "actor_id", "actor_name"] + SCHEMA["evidence"][2:]
    ev = sorted(db["evidence"], key=lambda e: e.get("date", ""), reverse=True)
    sheet("Evidence", ecols,
          [[actors.get(e["actor_id"], {}).get("name", "") if c == "actor_name" else e.get(c, "") for c in ecols]
           for e in ev],
          widths={"quote_original": 70, "quote_en": 70, "grading_note": 50, "source_url": 40},
          url_cols=("source_url", "archive_url"))
    gcols = SCHEMA["geo_units"][:5] + ["parent_name", "supporters"] + SCHEMA["geo_units"][6:]
    sheet("Local coverage", gcols,
          [[geo_name.get(x["parent_id"], "") if c == "parent_name" else
            geo_count.get(x["geo_unit_id"], 0) if c == "supporters" else x.get(c, "") for c in gcols]
           for x in db["geo_units"]], url_cols=("website",))
    lcols = ["link_id", "from_actor_id", "from_name", "link_type", "to_actor_id", "to_name", "evidence_id",
             "source_url", "note"]
    sheet("Links", lcols,
          [[actors.get(l["from_actor_id"], {}).get("name", "") if c == "from_name" else
            actors.get(l["to_actor_id"], {}).get("name", "") if c == "to_name" else l.get(c, "") for c in lcols]
           for l in db["links"]], url_cols=("source_url",))
    pcols = ["profile_id", "actor_id", "name", "platform", "url", "source", "checked"]
    sheet("Profiles", pcols,
          [[actors.get(p["actor_id"], {}).get("name", "") if c == "name" else p.get(c, "") for c in pcols]
           for p in db["profiles"]], url_cols=("url",))
    sheet("Harvest lists", SCHEMA["harvest"], [[h.get(c, "") for c in SCHEMA["harvest"]] for h in db["harvest"]],
          url_cols=("url",))
    sheet("Run log", SCHEMA["runs"], [[r.get(c, "") for c in SCHEMA["runs"]] for r in db["runs"]])
    former = [r for r in d if r["support_status"] == "retracted"]
    if former:  # omit the sheet while nobody has retracted
        sheet("Former supporters", cols, [[r.get(c, "") for c in cols] for r in former],
              url_cols=("strongest_source",))
    wb.save(out)
    print(f"Wrote {out}")


def cmd_merge(a):
    """Merge duplicate actor DROP into KEEP: move evidence, links and profiles, fill blanks, delete DROP."""
    need_dir(a.dir)
    db = load_all(a.dir)
    by = {r["actor_id"]: r for r in db["actors"]}
    if a.keep not in by or a.drop not in by or a.keep == a.drop:
        sys.exit("keep and drop must be two different existing actor IDs")
    keep, drop = by[a.keep], by[a.drop]
    for col in SCHEMA["actors"]:
        if col in ("actor_id", "notes", "updated") or keep.get(col):
            continue
        if drop.get(col):
            keep[col] = drop[col]
    extra = [x for x in (drop.get("other_types") or "").split(";") if x and x != keep["actor_type"]]
    if drop["actor_type"] != keep["actor_type"]:
        extra.append(drop["actor_type"])
    mine = [x for x in (keep.get("other_types") or "").split(";") if x]
    keep["other_types"] = ";".join(dict.fromkeys(mine + extra))
    if norm_name(drop["name"]) != norm_name(keep["name"]):
        keep["notes"] = (keep.get("notes", "") + f" Also listed as '{drop['name']}'.").strip()
    n = 0
    for e in db["evidence"]:
        if e["actor_id"] == a.drop:
            e["actor_id"] = a.keep; n += 1
    for l in db["links"]:
        for c in ("from_actor_id", "to_actor_id"):
            if l.get(c) == a.drop:
                l[c] = a.keep
    for p in db["profiles"]:
        if p["actor_id"] == a.drop:
            p["actor_id"] = a.keep
    for r in db["actors"]:
        if r.get("affiliation_actor_id") == a.drop:
            r["affiliation_actor_id"] = a.keep
    db["actors"] = [r for r in db["actors"] if r["actor_id"] != a.drop]
    for t_ in ("actors", "evidence", "links", "profiles"):
        save(a.dir, t_, db[t_])
    print(f"Merged {a.drop} ({drop['name']}) into {a.keep} ({keep['name']}); moved {n} evidence rows")


def _plain(s, tags=True):
    import html as _html
    s = _html.unescape(re.sub(r"<[^>]+>", " ", s) if tags else s)
    s = unicodedata.normalize("NFKC", s).lower()
    for a_, b_ in (("‘", "'"), ("’", "'"), ("“", '"'), ("”", '"'), ("–", "-"), ("—", "-")):
        s = s.replace(a_, b_)
    return re.sub(r"\s+", " ", s).strip()


def _pdf_to_text(data):
    """Text of a PDF via poppler's pdftotext when installed, else via pypdf (pip install pypdf), else None."""
    import shutil, subprocess, tempfile
    if shutil.which("pdftotext"):
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            f.write(data)
            path = f.name
        try:
            return subprocess.run(["pdftotext", path, "-"], capture_output=True, text=True).stdout
        finally:
            os.unlink(path)
    try:
        import io, pypdf
        return "\n".join(pg.extract_text() or "" for pg in pypdf.PdfReader(io.BytesIO(data)).pages)
    except Exception:
        return None


def _decode_page(raw, header_charset=None):
    """Decode a page in the charset it declares (header, then <meta>), else UTF-8, else the common East Asian ones.
    Strict decoding, so a wrong guess fails over instead of producing mojibake; Latin-1 is the last resort. A page that
    declares nothing and is not UTF-8 is tried as GB18030 first, which accepts most byte strings: such a page may need
    its charset added by hand (rare; nearly every page declares one)."""
    declared = header_charset
    if not declared:
        m = re.search(rb'charset\s*=\s*["\']?([A-Za-z0-9_-]+)', raw[:4096])
        declared = m.group(1).decode("ascii", "ignore") if m else None
    for enc in (declared, "utf-8", "gb18030", "cp932", "euc-kr", "latin-1"):
        if enc:
            try:
                return raw.decode(enc)
            except (UnicodeDecodeError, LookupError):
                continue
    return raw.decode("utf-8", "ignore")


def _kokkai_text(u, timeout):
    """Japan's Diet records (kokkai.ndl.go.jp/txt/<issue>/<speech>) are JavaScript pages; the same text is in the
    open API. Returns the speeches as one string, or None when u is not such a link."""
    import json, urllib.request
    m = re.search(r"kokkai\.ndl\.go\.jp/txt/([^/?#]+)(?:/(\d+))?", u)
    if not m:
        return None
    api = f"https://kokkai.ndl.go.jp/api/speech?recordPacking=json&issueID={m.group(1)}" + (f"&speechNumber={m.group(2)}" if m.group(2) else "")
    req = urllib.request.Request(api, headers={"User-Agent": "Mozilla/5.0 (wtdb quote check)"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return " ".join(r.get("speech", "") for r in data.get("speechRecord", []))


def _browser_text(url):
    """Fallback for pages a plain client cannot read (403, JavaScript): skill/scripts/fetch_page.py, if playwright is installed."""
    import importlib.util, subprocess
    helper = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fetch_page.py")
    if not os.path.exists(helper) or importlib.util.find_spec("playwright") is None:
        return None
    try:
        r = subprocess.run([sys.executable, helper, url, "--text", "--no-extract"], capture_output=True, text=True, timeout=180)
    except Exception:
        return None
    return r.stdout if r.returncode == 0 and r.stdout.strip() else None


def _fetch_for_check(u, timeout):
    """Page or PDF text for a quote check: plain HTTP first, then pdftotext for PDFs, then the headless browser."""
    import urllib.request
    raw = None
    try:
        text = _kokkai_text(u, timeout) if "kokkai.ndl.go.jp/txt/" in u else None
        if text:
            return _plain(text, tags=False)
        req = urllib.request.Request(u, headers={"User-Agent": "Mozilla/5.0 (wtdb quote check)"})
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw = resp.read()
            ctype = resp.headers.get("Content-Type", "").lower()
            charset = resp.headers.get_content_charset()
        if raw[:4] == b"%PDF" or "pdf" in ctype:
            text = _pdf_to_text(raw)
            return _plain(text, tags=False) if text else None  # a "<" in PDF text must not swallow what follows
        return _plain(_decode_page(raw, charset))
    except Exception:
        text = _browser_text(u)
        return _plain(text) if text else None


def _dehyphenate(page):
    """Line-end hyphens in PDFs split words ("miljardärs- skatt"); join them for a second comparison."""
    return re.sub(r"(?<=\w)- (?=[a-zåäö])", "", page)


CJK = re.compile(r"[\u3040-\u30ff\u3400-\u9fff\uac00-\ud7af]")


def _quote_parts(e):
    """The fragments of a quote to look for on its page (the quote split at ellipses), or [] when it is too short to
    check. Words are counted by spaces, and in Chinese, Japanese and Korean, which have none, by characters: a
    fragment needs three words or six CJK characters. A quote shorter than that is checked whole when it is a name
    on a roll-call or roster (two words, a vote row, or three CJK characters)."""
    q = _plain(e["quote_original"])
    frags = [p.strip() for p in re.split(r"…|\.\.\.", q) if p.strip()]
    parts = [p for p in frags if len(p.split()) >= 3 or (CJK.search(p) and len(p.strip()) >= 6)]
    if parts:
        return parts
    # name...organisation...title: every fragment is short, but each must still be on the page
    short = [p for p in frags if len(p.split()) >= 2 or (CJK.search(p) and len(p.replace(" ", "")) >= 3)]
    if len(frags) > 1 and len(short) == len(frags):
        return frags
    if len(q.split()) >= 2 or e.get("evidence_type") == "vote" or (CJK.search(q) and len(q) >= 3):
        return [q]
    return []


def cmd_checkquotes(a):
    """Fetch each source page (HTML, PDF via pdftotext, walled pages via the headless browser if installed) and check that the quote appears on it."""
    import random
    need_dir(a.dir)
    rows = [e for e in load(a.dir, "evidence") if e.get("source_url")]
    if a.run:
        rows = [e for e in rows if e.get("run_id") == a.run]
    if a.ids:
        rows = [e for e in rows if e["evidence_id"] in a.ids.split(",")]
    if a.sample and len(rows) > a.sample:
        random.seed(a.seed)
        rows = random.sample(rows, a.sample)
    cache, browser_cache, ok, miss, skip, short = {}, {}, 0, [], [], []
    for e in rows:
        u = e["source_url"]
        parts = _quote_parts(e)
        if not parts:
            short.append(e["evidence_id"]); continue
        if u not in cache:
            cache[u] = _fetch_for_check(u, a.timeout)
        page = cache[u]
        if page is None:
            skip.append(e["evidence_id"]); continue

        def found(pg):
            return bool(parts) and all(any(p.strip(" .") in v for v in (pg, _dehyphenate(pg))) for p in parts)

        if not found(page):
            # some sites serve a consent or shortened variant to plain clients: try the headless browser before calling it a miss
            if u not in browser_cache:
                bt = _browser_text(u)
                browser_cache[u] = _plain(bt) if bt else None
            page = browser_cache[u] or page
        if found(page):
            ok += 1
        else:
            miss.append(e["evidence_id"])
    print(f"Checked {len(rows)} rows: {ok} found on the page, {len(miss)} NOT found, {len(skip)} skipped (unreadable: no pdftotext, blocked or error), "
          f"{len(short)} too short to check")
    if miss:
        print("Not found:", ", ".join(miss))
    if short:
        print("Too short to check:", ", ".join(short[:40]) + (" ..." if len(short) > 40 else ""))
    if skip:
        print("Skipped:", ", ".join(skip[:40]) + (" ..." if len(skip) > 40 else ""))
    if miss and a.strict:
        sys.exit(1)


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    s = p.add_subparsers(dest="cmd", required=True)
    x = s.add_parser("init"); x.add_argument("dir"); x.add_argument("--area", required=True); x.set_defaults(f=cmd_init)
    x = s.add_parser("add"); x.add_argument("dir"); x.add_argument("table"); x.add_argument("rows"); x.set_defaults(f=cmd_add)
    x = s.add_parser("set"); x.add_argument("dir"); x.add_argument("table"); x.add_argument("id")
    x.add_argument("pairs", nargs="+"); x.set_defaults(f=cmd_set)
    for name, fn in (("validate", cmd_validate), ("dedupe", cmd_dedupe), ("status", cmd_status)):
        x = s.add_parser(name); x.add_argument("dir"); x.set_defaults(f=fn)
    x = s.add_parser("locate"); x.add_argument("dir"); x.set_defaults(f=cmd_locate)
    x = s.add_parser("merge"); x.add_argument("dir"); x.add_argument("keep"); x.add_argument("drop"); x.set_defaults(f=cmd_merge)
    x = s.add_parser("checkquotes"); x.add_argument("dir"); x.add_argument("--run"); x.add_argument("--ids")
    x.add_argument("--sample", type=int, default=0); x.add_argument("--seed", type=int, default=1)
    x.add_argument("--timeout", type=int, default=25); x.add_argument("--strict", action="store_true")
    x.set_defaults(f=cmd_checkquotes)
    x = s.add_parser("export"); x.add_argument("dir"); x.add_argument("--out"); x.set_defaults(f=cmd_export)
    x = s.add_parser("schema"); x.add_argument("--write", action="store_true"); x.set_defaults(f=cmd_schema)
    a = p.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
